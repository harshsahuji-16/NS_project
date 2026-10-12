"""Generate professional 6-slide PowerPoint presentation for Classical Cryptography Toolkit."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def create_presentation():
    os.makedirs("docs", exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    DARK_BG = RGBColor(11, 15, 25)
    CARD_BG = RGBColor(17, 24, 39)
    CYAN_ACCENT = RGBColor(6, 182, 212)
    GREEN_ACCENT = RGBColor(16, 185, 129)
    TEXT_LIGHT = RGBColor(243, 244, 246)
    TEXT_MUTED = RGBColor(156, 163, 175)

    def set_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = DARK_BG

    def add_header(slide, title_text, category_text="NETWORK SECURITY MINI PROJECT"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN_ACCENT

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(2.2))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "DESIGN AND DEVELOPMENT OF A\nCLASSICAL CRYPTOGRAPHY TOOLKIT"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = CYAN_ACCENT

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "An Interactive Educational Platform for Classical Encryption and Decryption"
    p1_sub.font.size = Pt(18)
    p1_sub.font.color.rgb = GREEN_ACCENT

    # Student metadata card
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.2), Inches(11.3), Inches(2.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CYAN_ACCENT
    card.line.width = Pt(1.5)

    cbox = s1.shapes.add_textbox(Inches(1.3), Inches(4.4), Inches(10.7), Inches(2.1))
    ctf = cbox.text_frame
    cp = ctf.paragraphs[0]
    cp.text = "ACADEMIC SUBMISSION DETAILS"
    cp.font.bold = True
    cp.font.size = Pt(13)
    cp.font.color.rgb = CYAN_ACCENT

    details = [
        "Course: Network Security",
        "Class: Final Year (ECM), Academic Year 2026–27",
        "Student Name: [ENTER YOUR NAME] | Roll No: [ENTER YOUR ROLL NUMBER]",
        "Batch: [ENTER YOUR BATCH] | Submission Date: [ENTER SUBMISSION DATE]",
    ]
    for d in details:
        p = ctf.add_paragraph()
        p.text = "• " + d
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 2: OBJECTIVES & MOTIVATION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Objectives, Problem Statement & Motivation")

    # Column 1: Problem Statement
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CYAN_ACCENT

    b1 = s2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    tf_b1 = b1.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "PROBLEM STATEMENT & MOTIVATION"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = CYAN_ACCENT

    points1 = [
        "Abstract Cryptographic Concepts: Students frequently struggle to visualize the mathematical mechanisms of classical ciphers.",
        "Scattered & Incomplete Implementations: Most online tools lack mathematical rigor, handle edge cases poorly, or omit reversible decryption.",
        "Vulnerability Misconceptions: Learners often conflate algorithmic speed with cryptographic security.",
        "Educational Need: A unified, offline, interactive laboratory providing mathematically verified algorithms with transparent execution steps.",
    ]
    for pt in points1:
        p = tf_b1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    # Column 2: Objectives
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = GREEN_ACCENT

    b2 = s2.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_b2 = b2.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "PROJECT OBJECTIVES"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = GREEN_ACCENT

    points2 = [
        "Implement 8 Classical Ciphers: Full mathematical implementation of 6 substitution and 2 transposition algorithms.",
        "Lossless Round-Trip Decryption: Guarantee that valid inputs and keys perfectly reconstruct the original message.",
        "Interactive Cybersecurity UI: Build a dark-themed Streamlit dashboard with instant step traces and matrix visualizers.",
        "High-Precision Benchmarking: Compare algorithm throughput and execution time using Python's time.perf_counter().",
        "Lossless File Processing: Provide secure text-file encryption and decryption with embedded metadata.",
    ]
    for pt in points2:
        p = tf_b2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 3: SYSTEM ARCHITECTURE & FLOWCHART
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "System Architecture & Operational Workflow")

    card3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    card3.fill.solid()
    card3.fill.fore_color.rgb = CARD_BG
    card3.line.color.rgb = CYAN_ACCENT

    b3 = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
    tf_b3 = b3.text_frame
    tf_b3.word_wrap = True
    p = tf_b3.paragraphs[0]
    p.text = "MODULAR LAYERED ARCHITECTURE"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = CYAN_ACCENT

    arch_points = [
        "Presentation Layer (Streamlit): Cyber-dark interface with sidebar navigation, 8 views, interactive widgets, step breakdowns, and Matplotlib timing graphs.",
        "Validation Layer (modules/validation.py): Centralized gatekeeper enforcing coprime determinants mod 26, 26-letter unique alphabets, keyword formats, and rail boundaries.",
        "Cryptographic Engine (crypto/): Pure mathematical Python implementations of Caesar, Monoalphabetic, Playfair, Hill, Vigenère, OTP, Rail Fence, and Columnar Transposition.",
        "File Management Layer (modules/file_handler.py): Handles UTF-8 text file packaging with structured metadata headers (# CIPHER, # TIMESTAMP, # METADATA) for seamless round-trips.",
        "Performance Benchmarking (modules/comparison.py): Microsecond benchmarking utilizing time.perf_counter() over multiple runs with round-trip verification.",
        "Testing Suite (tests/): 51 automated unit tests executed via pytest verifying edge cases, boundary conditions, and textbook test vectors.",
    ]
    for ap in arch_points:
        p = tf_b3.add_paragraph()
        p.text = "✔ " + ap
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 4: THE EIGHT ALGORITHMS
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "The Eight Implemented Algorithms")

    # Substitution Box
    c_sub = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c_sub.fill.solid()
    c_sub.fill.fore_color.rgb = CARD_BG
    c_sub.line.color.rgb = CYAN_ACCENT

    t_sub = s4.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_sub = t_sub.text_frame
    tf_sub.word_wrap = True
    p = tf_sub.paragraphs[0]
    p.text = "SUBSTITUTION CIPHERS (6)"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = CYAN_ACCENT

    sub_algos = [
        "1. Caesar Cipher: C = (P + k) mod 26. Shift of 5: 'NETWORK' -> 'SJYBTWP'.",
        "2. Monoalphabetic: 26! keyspace (~4x10^26). Preserves case, vulnerable to frequency analysis.",
        "3. Playfair: 5x5 key matrix. Digraph substitution (Row/Col/Rect). Combines I/J, resolves double letters with 'X'.",
        "4. Hill Cipher: 2x2 matrix mod 26. Coprime det gcd(det, 26)=1. Matrix mult C = (K·P) mod 26.",
        "5. Vigenère: Polyalphabetic periodic key. Preserves casing & punctuation without advancing key stream.",
        "6. One-Time Pad: Perfect secrecy if key is truly random, secret, equal length, and used exactly once.",
    ]
    for sa in sub_algos:
        p = tf_sub.add_paragraph()
        p.text = "• " + sa
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    # Transposition Box
    c_tra = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c_tra.fill.solid()
    c_tra.fill.fore_color.rgb = CARD_BG
    c_tra.line.color.rgb = GREEN_ACCENT

    t_tra = s4.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_tra = t_tra.text_frame
    tf_tra.word_wrap = True
    p = tf_tra.paragraphs[0]
    p.text = "TRANSPOSITION CIPHERS (2)"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = GREEN_ACCENT

    tra_algos = [
        "7. Rail Fence (Zigzag): Characters written in diagonal zigzag across d rails and read horizontally. Preserves all characters and spacing.",
        "8. Columnar Transposition: Plaintext written into rectangle of width k; columns read in alphabetical order of key.",
        "Tie-Breaking Rule: Duplicate key letters are ordered left-to-right via stable sorting.",
        "Irregular Rows: Handles incomplete rows dynamically without mandatory padding, or optionally pads with stored metadata.",
    ]
    for ta in tra_algos:
        p = tf_tra.add_paragraph()
        p.text = "• " + ta
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 5: VERIFICATION & TESTING
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Verification, Testing & Sample Results")

    c5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    c5.fill.solid()
    c5.fill.fore_color.rgb = CARD_BG
    c5.line.color.rgb = CYAN_ACCENT

    t5 = s5.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
    tf5 = t5.text_frame
    tf5.word_wrap = True
    p = tf5.paragraphs[0]
    p.text = "COMPREHENSIVE AUTOMATED TESTING (51 / 51 PASSED)"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = GREEN_ACCENT

    test_points = [
        "Caesar Reference Vector: 'NETWORK SECURITY' with key 5 -> 'SJYBTWP XJHZWNYD' -> Decrypted: 'NETWORK SECURITY'. [PASSED]",
        "Vigenère Reference Vector: 'ATTACKATDAWN' with 'LEMON' -> 'LXFOPVEFRNHR' -> Decrypted: 'ATTACKATDAWN'. [PASSED]",
        "Hill Matrix Validation: Accepted [[3,3],[2,5]] (det=9). Rejected singular det=0, even det=2, factor-13 det=13. [PASSED]",
        "Hill Textbook Vector: 'HELP' with [[3,3],[2,5]] -> 'HIAT' -> Decrypted: 'HELP'. [PASSED]",
        "Playfair Digraphs: 'BALLOON' with 'MONARCHY' -> correct filler 'X' insertion -> digraphs restored. [PASSED]",
        "Rail Fence Zigzag: 'DEFEND THE EAST WALL' with 3 rails -> round-trip restored with boundary handling. [PASSED]",
        "Columnar Transposition: 'DEFEND THE EAST WALL AT DAWN!' with 'GERMAN' -> lossless irregular recovery. [PASSED]",
        "File Round-Trip: UTF-8 structured packaging with metadata headers -> flawless recovery of multi-line files. [PASSED]",
    ]
    for tp in test_points:
        p = tf5.add_paragraph()
        p.text = "✔ " + tp
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 6: RESULTS, LIMITATIONS & CONCLUSION
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Results, Limitations, Future Scope & Conclusion")

    c6_1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c6_1.fill.solid()
    c6_1.fill.fore_color.rgb = CARD_BG
    c6_1.line.color.rgb = CYAN_ACCENT

    t6_1 = s6.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf6_1 = t6_1.text_frame
    tf6_1.word_wrap = True
    p = tf6_1.paragraphs[0]
    p.text = "LIMITATIONS & SECURITY CRITIQUE"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = CYAN_ACCENT

    lims = [
        "Vulnerability to Frequency Analysis: Monoalphabetic and Caesar preserve single-letter frequency signatures of natural language.",
        "Small Key Spaces: Caesar (25 keys) can be brute-forced instantly; Hill 2x2 matrix can be recovered with known plaintext pairs.",
        "Lack of Authentication / Integrity: Classical ciphers do not protect against tampering, bit-flipping, or malleability.",
        "Not Suitable for Modern Security: Must never be deployed for confidential communications. Modern security requires AES-256-GCM, RSA-4096, and ECC.",
    ]
    for lm in lims:
        p = tf6_1.add_paragraph()
        p.text = "• " + lm
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    c6_2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c6_2.fill.solid()
    c6_2.fill.fore_color.rgb = CARD_BG
    c6_2.line.color.rgb = GREEN_ACCENT

    t6_2 = s6.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf6_2 = t6_2.text_frame
    tf6_2.word_wrap = True
    p = tf6_2.paragraphs[0]
    p.text = "CONCLUSION & FUTURE SCOPE"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = GREEN_ACCENT

    concs = [
        "Successful Delivery: Built a production-grade, menu-driven toolkit with 8 classical algorithms, 100% passing tests, and modern cyber UI.",
        "Educational Value: Demystifies historical cipher mathematics through real-time matrix, digraph, and zigzag step traces.",
        "Future Enhancements: Integrate automated cryptanalysis modules (Index of Coincidence, Kasiski examination, n-gram frequency solvers).",
        "Extended Support: Implement mechanical cipher simulation (Enigma M3/M4 machine) and 3x3 Hill matrices.",
    ]
    for cn in concs:
        p = tf6_2.add_paragraph()
        p.text = "• " + cn
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    output_path = os.path.join("docs", "Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")


if __name__ == "__main__":
    create_presentation()
